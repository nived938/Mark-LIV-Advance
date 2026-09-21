"""Local DOCX/XLSX/PPTX/PDF generator for JARVIS."""

from __future__ import annotations

import json
import re
from pathlib import Path

OUT_DIR = Path.home() / "Documents" / "JARVIS Office"


def _safe_name(value: str, default: str = "jarvis_document") -> str:
    clean = re.sub(r"[^A-Za-z0-9._ -]+", "_", str(value or "").strip())
    clean = re.sub(r"\s+", " ", clean).strip(" .")
    return (clean or default)[:120]


def _output_path(filename: str, kind: str) -> Path:
    ext = {"docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx", "pdf": ".pdf"}[kind]
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    raw = str(filename or "").strip()
    path = Path(raw).expanduser() if raw else OUT_DIR / f"jarvis_{kind}"
    if not path.is_absolute():
        path = OUT_DIR / path
    if path.suffix.lower() != ext:
        path = path.with_suffix(ext)
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _json(value, default):
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(str(value or ""))
    except Exception:
        return default


def _build_docx(path: Path, title: str, content: str, sections) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = title
    doc.add_heading(title, level=0)
    blocks = sections or [{"heading": "", "body": content}]
    for block in blocks:
        if not isinstance(block, dict):
            continue
        heading = str(block.get("heading") or "").strip()
        body = str(block.get("body") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        for line in body.splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith(("- ", "* ")):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(11)
    doc.save(path)


def _build_xlsx(path: Path, title: str, content: str, rows) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    data = _json(rows, [])
    if not isinstance(data, list) or not data:
        data = [[v.strip() for v in line.split("|")] for line in content.splitlines() if line.strip()]

    wb = Workbook()
    ws = wb.active
    ws.title = _safe_name(title, "Sheet")[:31]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    width = max(1, max((len(r) for r in data if isinstance(r, list)), default=1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)

    for r_idx, row in enumerate(data, 3):
        if not isinstance(row, list):
            continue
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, str(value)).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    for column_cells in ws.columns:
        max_len = min(
            50,
            max((len(str(cell.value or "")) for cell in column_cells), default=8),
        )
        ws.column_dimensions[column_cells[0].column_letter].width = max(10, max_len + 2)
    wb.save(path)


def _build_pptx(path: Path, title: str, content: str, slides) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    data = _json(slides, [])
    if not isinstance(data, list) or not data:
        data = [{"title": title, "bullets": [x.strip() for x in content.splitlines() if x.strip()]}]

    prs = Presentation()
    for item in data:
        if not isinstance(item, dict):
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title") or title)
        bullets = item.get("bullets") or item.get("body") or []
        if isinstance(bullets, str):
            bullets = [x.strip() for x in bullets.splitlines() if x.strip()]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for i, bullet in enumerate(bullets):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = str(bullet)
            para.font.size = Pt(20)
    if not prs.slides:
        prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(path)


def _build_pdf(path: Path, title: str, content: str, sections) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()
    title_style = styles["Title"].clone("JarvisTitle")
    title_style.alignment = TA_CENTER
    story = [Paragraph(escape(title), title_style), Spacer(1, 16)]
    blocks = sections or [{"heading": "", "body": content}]
    for block in blocks:
        if not isinstance(block, dict):
            continue
        heading = str(block.get("heading") or "").strip()
        body = str(block.get("body") or "").strip()
        if heading:
            story.extend([Paragraph(escape(heading), styles["Heading2"]), Spacer(1, 6)])
        for line in body.splitlines():
            if line.strip():
                story.extend([Paragraph(escape(line.strip()), styles["BodyText"]), Spacer(1, 5)])
            else:
                story.append(Spacer(1, 6))
    SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)


TOOL = {
    "name": "office_builder",
    "description": (
        "Generate office documents locally as DOCX, XLSX, PPTX, or PDF. "
        "Use title/content for simple documents, or sections/slides/rows as JSON "
        "strings for structured output. Files are saved in Documents/JARVIS Office by default."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "kind": {"type": "STRING", "description": "docx | xlsx | pptx | pdf"},
            "title": {"type": "STRING", "description": "Document title."},
            "content": {"type": "STRING", "description": "Body text."},
            "sections": {"type": "STRING", "description": "JSON array of {heading,body} for DOCX/PDF."},
            "slides": {"type": "STRING", "description": "JSON array of {title,bullets} for PPTX."},
            "rows": {"type": "STRING", "description": "JSON array of row arrays for XLSX."},
            "filename": {"type": "STRING", "description": "Optional output filename/path."},
        },
        "required": ["kind", "title"],
    },
}


def office_builder(parameters: dict | None = None, **_) -> str:
    p = parameters or {}
    kind = str(p.get("kind") or "").strip().lower()
    title = str(p.get("title") or "JARVIS Document").strip()
    content = str(p.get("content") or "").strip()
    if kind not in {"docx", "xlsx", "pptx", "pdf"}:
        return "Kind must be docx, xlsx, pptx, or pdf."
    path = _output_path(p.get("filename", ""), kind)
    try:
        if kind == "docx":
            _build_docx(path, title, content, _json(p.get("sections"), []))
        elif kind == "xlsx":
            _build_xlsx(path, title, content, p.get("rows", ""))
        elif kind == "pptx":
            _build_pptx(path, title, content, p.get("slides", ""))
        else:
            _build_pdf(path, title, content, _json(p.get("sections"), []))
        return f"Created {kind.upper()}: {path}"
    except ImportError as exc:
        return f"{kind.upper()} generator dependency is missing: {exc}"
    except Exception as exc:
        return f"Could not create {kind.upper()}: {exc}"
