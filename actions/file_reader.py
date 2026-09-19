"""
file_reader.py

Reads files and returns their contents to Mark-LIV.

Supported:
- TXT
- Markdown
- Python
- JavaScript
- TypeScript
- HTML
- CSS
- JSON
- XML
- YAML
- CSV
- DOCX
- PDF
- XLSX
- PPTX
"""

import os
import json
import platform
from pathlib import Path


_SYSTEM = platform.system()


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",

    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",

    ".html",
    ".htm",
    ".css",

    ".json",
    ".xml",

    ".yaml",
    ".yml",

    ".csv",
    ".tsv",

    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".kts",

    ".sh",
    ".bash",
    ".ps1",

    ".sql",
}


def _find_file(file_path: str) -> Path | None:

    if not file_path:
        return None

    file_path = str(file_path).strip().strip('"').strip("'")

    # Expand environment variables such as %USERPROFILE%
    file_path = os.path.expandvars(file_path)

    # Expand ~
    file_path = os.path.expanduser(file_path)

    path = Path(file_path)

    # Direct path
    if path.exists():
        return path.resolve()

    # Search relative to current directory
    current = Path.cwd() / file_path

    if current.exists():
        return current.resolve()

    # Windows common locations
    if _SYSTEM == "Windows":

        home = Path.home()

        common_locations = [
            home / "Desktop",
            home / "Documents",
            home / "Downloads",
            home / "Pictures",
            home / "Videos",
            home / "Music",
        ]

        for location in common_locations:

            candidate = location / file_path

            if candidate.exists():
                return candidate.resolve()

            # Also search by filename
            filename_candidate = location / Path(file_path).name

            if filename_candidate.exists():
                return filename_candidate.resolve()

    return None


def _read_text(path: Path) -> str:

    try:

        return path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

    except Exception as e:

        return f"Could not read text file: {e}"


def _read_docx(path: Path) -> str:

    try:

        from docx import Document

        document = Document(path)

        parts = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                parts.append(paragraph.text)

        # Read tables too
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:
                    row_text.append(cell.text)

                parts.append(" | ".join(row_text))

        return "\n".join(parts)

    except ImportError:

        return (
            "python-docx is not installed. "
            "Run: python -m pip install python-docx"
        )

    except Exception as e:

        return f"Could not read DOCX file: {e}"


def _read_pdf(path: Path) -> str:

    # Try pdfplumber first
    try:

        import pdfplumber

        parts = []

        with pdfplumber.open(path) as pdf:

            for number, page in enumerate(pdf.pages, start=1):

                text = page.extract_text()

                if text:

                    parts.append(
                        f"--- Page {number} ---\n{text}"
                    )

        if parts:
            return "\n\n".join(parts)

    except ImportError:
        pass

    except Exception as e:
        print(f"[file_reader] pdfplumber failed: {e}")

    # Fallback to PyPDF2
    try:

        import PyPDF2

        parts = []

        with open(path, "rb") as file:

            reader = PyPDF2.PdfReader(file)

            for number, page in enumerate(
                reader.pages,
                start=1
            ):

                text = page.extract_text() or ""

                if text:

                    parts.append(
                        f"--- Page {number} ---\n{text}"
                    )

        if parts:
            return "\n\n".join(parts)

    except ImportError:

        return (
            "PDF reading requires pdfplumber or PyPDF2. "
            "Run: python -m pip install pdfplumber PyPDF2"
        )

    except Exception as e:

        return f"Could not read PDF: {e}"

    return (
        "No readable text was found in this PDF. "
        "It may be a scanned/image-only PDF."
    )


def _read_xlsx(path: Path) -> str:

    try:

        import openpyxl

        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        output = []

        for sheet in workbook.worksheets:

            output.append(
                f"--- Sheet: {sheet.title} ---"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = []

                for value in row:

                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value))

                if any(values):

                    output.append(
                        " | ".join(values)
                    )

        workbook.close()

        return "\n".join(output)

    except ImportError:

        return (
            "openpyxl is not installed. "
            "Run: python -m pip install openpyxl"
        )

    except Exception as e:

        return f"Could not read Excel file: {e}"


def _read_pptx(path: Path) -> str:

    try:

        from pptx import Presentation

        presentation = Presentation(path)

        output = []

        for number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            output.append(
                f"--- Slide {number} ---"
            )

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        output.append(text)

        return "\n".join(output)

    except ImportError:

        return (
            "python-pptx is not installed. "
            "Run: python -m pip install python-pptx"
        )

    except Exception as e:

        return f"Could not read PowerPoint file: {e}"


def _read_csv(path: Path) -> str:

    try:

        import csv

        output = []

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore",
            newline="",
        ) as file:

            reader = csv.reader(file)

            for row in reader:

                output.append(
                    " | ".join(row)
                )

        return "\n".join(output)

    except Exception as e:

        return f"Could not read CSV file: {e}"


def _read_file(path: Path) -> str:

    extension = path.suffix.lower()

    if extension == ".docx":
        return _read_docx(path)

    if extension == ".pdf":
        return _read_pdf(path)

    if extension in (
        ".xlsx",
        ".xlsm",
    ):
        return _read_xlsx(path)

    if extension == ".pptx":
        return _read_pptx(path)

    if extension in (
        ".csv",
        ".tsv",
    ):
        return _read_csv(path)

    if extension in TEXT_EXTENSIONS:
        return _read_text(path)

    # Last attempt for unknown files
    try:

        return _read_text(path)

    except Exception:

        return (
            f"I don't know how to read this file type: "
            f"{extension}"
        )


def read_file(
    parameters=None,
    response=None,
    player=None,
    session_memory=None,
) -> str:

    parameters = parameters or {}

    file_path = (
        parameters.get("file_path")
        or parameters.get("path")
        or parameters.get("file")
        or parameters.get("filename")
        or ""
    )

    file_path = str(file_path).strip()

    if not file_path:

        return (
            "Please provide the file path or filename. "
            "For example: read_file('C:\\\\Users\\\\USER\\\\Downloads\\\\test.txt')"
        )

    path = _find_file(file_path)

    if path is None:

        return (
            f"I could not find the file '{file_path}'. "
            "Please provide the full path or filename."
        )

    if not path.is_file():

        return f"'{path}' is not a file."

    try:

        size = path.stat().st_size

    except Exception:

        size = 0

    # Prevent accidentally sending enormous files
    max_chars = int(
        parameters.get(
            "max_chars",
            50000
        )
    )

    content = _read_file(path)

    if content.startswith(
        (
            "Could not ",
            "No readable ",
            "I don't know ",
            "python-",
            "openpyxl ",
            "PDF reading ",
        )
    ):
        return content

    if len(content) > max_chars:

        content = content[:max_chars]

        content += (
            "\n\n[File content truncated. "
            f"Only the first {max_chars:,} characters were read.]"
        )

    extension = path.suffix.lower()

    return (
        f"File: {path.name}\n"
        f"Path: {path}\n"
        f"Type: {extension or 'unknown'}\n"
        f"Size: {size:,} bytes\n\n"
        f"Content:\n"
        f"{content}"
    )


TOOL = {
    "name": "read_file",
    "description": (
        "Reads the contents of a file and returns the content to Mark-LIV. "
        "Use this when the user asks to read, inspect, understand, summarize, "
        "or show the contents of a file. Supports text files, source code, "
        "JSON, CSV, Markdown, DOCX, PDF, Excel XLSX, and PowerPoint PPTX. "
        "The user should provide a filename or path."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "file_path": {
                "type": "STRING",
                "description": (
                    "Path or filename of the file to read. "
                    "Example: C:\\Users\\USER\\Downloads\\report.pdf"
                )
            },
            "max_chars": {
                "type": "INTEGER",
                "description": (
                    "Maximum number of characters to return. "
                    "Default is 50000."
                )
            }
        },
        "required": [
            "file_path"
        ]
    },
    "handler": read_file,
}